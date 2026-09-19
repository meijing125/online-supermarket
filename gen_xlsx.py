# -*- coding: utf-8 -*-
"""解析 商城_测试用例.md，生成格式化 Excel"""
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = '商城_测试用例.md'
DST = '商城_测试用例.xlsx'

# ---------- 解析 Markdown ----------
lines = open(SRC, encoding='utf-8').read().splitlines()
sections = []  # {level, title, headers, rows}
cur = None
for line in lines:
    s = line.strip()
    m = re.match(r'^(#{2,3})\s+(.*)$', s)
    if m:
        cur = {'level': len(m.group(1)), 'title': m.group(2).strip(), 'headers': None, 'rows': []}
        sections.append(cur)
        continue
    if s.startswith('|') and cur is not None:
        cells = [c.strip() for c in s.strip('|').split('|')]
        if re.match(r'^[\s:\-|]+$', s):  # 分隔行
            continue
        if cur['headers'] is None:
            cur['headers'] = cells
        else:
            cur['rows'].append(cells)

# ---------- 样式 ----------
HEADER_FILL = PatternFill('solid', fgColor='1F4E79')
HEADER_FONT = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
SEC_FILL   = PatternFill('solid', fgColor='D9E2F3')
SEC_FONT   = Font(name='微软雅黑', size=11, bold=True, color='1F4E79')
BODY_FONT  = Font(name='微软雅黑', size=10)
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)

# 模块 → 优先级
PRIO = {
    '2.1': 'P0', '2.2': 'P0', '2.5': 'P0', '2.6': 'P0', '2.7': 'P0',
    '2.8': 'P1', '2.9': 'P1', '2.10': 'P1', '2.11': 'P1', '2.12': 'P1',
    '2.3': 'P2', '2.4': 'P2', '2.13': 'P2',
}


def get_sec(prefix):
    for sec in sections:
        if sec['title'].startswith(prefix):
            return sec
    return None


def write_table(ws, headers, rows, widths, start_row, prio_col=None):
    """写表头 + 数据行，返回下一个可用行号"""
    # 表头
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    r = start_row + 1
    for row in rows:
        vals = list(row)
        if prio_col is not None:
            vals.append(prio_col)
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = WRAP
        r += 1
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    return r


def write_section_title(ws, title, ncols, row):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    cell = ws.cell(row=row, column=1, value=title)
    cell.fill = SEC_FILL
    cell.font = SEC_FONT
    cell.alignment = Alignment(vertical='center')
    ws.row_dimensions[row].height = 22
    return row + 1


wb = Workbook()

# ========== Sheet 1: 功能测试用例 ==========
ws1 = wb.active
ws1.title = '功能测试用例'
headers1 = ['编号', '模块', '测试项', '前置条件', '步骤', '预期结果', '优先级']
widths1 = [11, 11, 18, 22, 46, 48, 9]
row = 1
n1 = 0
for sec in sections:
    if sec['title'][:2] not in ('2.',):
        continue
    prio = PRIO.get(sec['title'][:3], 'P2')
    row = write_section_title(ws1, sec['title'], len(headers1), row)
    row = write_table(ws1, headers1, sec['rows'], widths1, row, prio)
    n1 += len(sec['rows'])
ws1.freeze_panes = 'A2'

# ========== Sheet 2: 安全与异常测试 ==========
ws2 = wb.create_sheet('安全与异常测试')
sec3 = get_sec('三、')
headers2 = ['编号', '模块', '测试项', '前置条件', '步骤', '预期结果']
widths2 = [11, 11, 18, 22, 46, 48]
row = write_section_title(ws2, sec3['title'], len(headers2), 1)
row = write_table(ws2, headers2, sec3['rows'], widths2, row)
n2 = len(sec3['rows'])
ws2.freeze_panes = 'A2'

# ========== Sheet 3: 疑似缺陷 ==========
ws3 = wb.create_sheet('疑似缺陷')
sec4 = get_sec('四、')
headers3 = ['编号', '严重级', '位置', '问题描述', '影响']
widths3 = [8, 9, 28, 58, 36]
row = write_section_title(ws3, sec4['title'], len(headers3), 1)
row = write_table(ws3, headers3, sec4['rows'], widths3, row)
n3 = len(sec4['rows'])
ws3.freeze_panes = 'A2'

# ========== Sheet 4: 测试说明 ==========
ws4 = wb.create_sheet('测试说明')
ws4.column_dimensions['A'].width = 18
ws4.column_dimensions['B'].width = 90
r = 1
info = [
    ('测试对象', '臻品汇数码奢品商城'),
    ('技术栈', 'Python Flask + MySQL(PyMySQL) + 原生 JavaScript + Bootstrap5'),
    ('启动方式', 'python app.py → http://localhost:5000'),
    ('默认管理员', 'admin / admin123'),
    ('用例规模', f'功能测试 {n1} 条 + 安全异常 {n2} 条 + 疑似缺陷 {n3} 条'),
]
for k, v in info:
    c1 = ws4.cell(row=r, column=1, value=k)
    c1.font = Font(name='微软雅黑', size=10, bold=True)
    c2 = ws4.cell(row=r, column=2, value=v)
    c2.font = BODY_FONT
    c2.alignment = WRAP
    r += 1

# 角色权限表
r += 1
r = write_section_title(ws4, '角色与权限', 2, r)
role = get_sec('1.1')
for row in role['rows']:
    for c, v in enumerate(row, 1):
        cell = ws4.cell(row=r, column=c, value=v)
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = WRAP
    r += 1

# 测试范围建议
r += 1
r = write_section_title(ws4, '测试范围建议', 2, r)
sec5 = get_sec('五、')
for row in sec5['rows']:
    for c, v in enumerate(row, 1):
        cell = ws4.cell(row=r, column=c, value=v)
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = WRAP
    r += 1

wb.save(DST)
print(f'生成成功: {DST}')
print(f'功能测试用例 {n1} 条, 安全异常 {n2} 条, 疑似缺陷 {n3} 条')
